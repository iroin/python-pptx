# encoding: utf-8

"""Test suite for pptx.text module."""

from __future__ import absolute_import

import pytest

from lxml import objectify

from hamcrest import assert_that, equal_to, is_, same_instance
from mock import MagicMock, Mock, patch

from pptx.constants import MSO, PP
from pptx.oxml import parse_xml_bytes
from pptx.oxml.core import SubElement
from pptx.oxml.ns import namespaces, nsdecls
from pptx.text import _Font, _Paragraph, _Run, TextFrame
from pptx.util import Pt

from .oxml.unitdata.text import an_rPr, test_text_objects, test_text_xml
from .unitutil import (
    absjoin, parse_xml_file, serialize_xml, TestCase, test_file_dir
)


nsmap = namespaces('a', 'r', 'p')


def actual_xml(elm):
    objectify.deannotate(elm, cleanup_namespaces=True)
    return serialize_xml(elm, pretty_print=True)


class Describe_Font(object):

    def it_knows_the_bold_setting(self, font, bold_font, bold_off_font):
        assert font.bold is None
        assert bold_font.bold is True
        assert bold_off_font.bold is False

    def it_can_change_the_bold_setting(
            self, font, bold_rPr_xml, bold_off_rPr_xml, rPr_xml):
        assert actual_xml(font._rPr) == rPr_xml
        font.bold = True
        assert actual_xml(font._rPr) == bold_rPr_xml
        font.bold = False
        assert actual_xml(font._rPr) == bold_off_rPr_xml
        font.bold = None
        assert actual_xml(font._rPr) == rPr_xml

    def it_can_set_the_font_size(self, font):
        font.size = 2400
        expected_xml = an_rPr().with_nsdecls().with_sz(2400).xml()
        assert actual_xml(font._rPr) == expected_xml

    # fixtures ---------------------------------------------

    @pytest.fixture
    def bold_font(self, bold_rPr):
        return _Font(bold_rPr)

    @pytest.fixture
    def bold_off_font(self, bold_off_rPr):
        return _Font(bold_off_rPr)

    @pytest.fixture
    def bold_off_rPr(self, bold_off_rPr_bldr):
        return bold_off_rPr_bldr.element

    @pytest.fixture
    def bold_off_rPr_bldr(self):
        return an_rPr().with_nsdecls().with_b(0)

    @pytest.fixture
    def bold_off_rPr_xml(self, bold_off_rPr_bldr):
        return bold_off_rPr_bldr.xml()

    @pytest.fixture
    def bold_rPr(self, bold_rPr_bldr):
        return bold_rPr_bldr.element

    @pytest.fixture
    def bold_rPr_bldr(self):
        return an_rPr().with_nsdecls().with_b(1)

    @pytest.fixture
    def bold_rPr_xml(self, bold_rPr_bldr):
        return bold_rPr_bldr.xml()

    @pytest.fixture
    def rPr(self, rPr_bldr):
        return rPr_bldr.element

    @pytest.fixture
    def rPr_bldr(self):
        return an_rPr().with_nsdecls()

    @pytest.fixture
    def rPr_xml(self, rPr_bldr):
        return rPr_bldr.xml()

    @pytest.fixture
    def font(self, rPr):
        return _Font(rPr)


class Describe_Paragraph(TestCase):

    def setUp(self):
        path = absjoin(test_file_dir, 'slide1.xml')
        self.sld = parse_xml_file(path).getroot()
        xpath = './p:cSld/p:spTree/p:sp/p:txBody/a:p'
        self.pList = self.sld.xpath(xpath, namespaces=nsmap)

        self.test_text = 'test text'
        self.p_xml = ('<a:p %s><a:r><a:t>%s</a:t></a:r></a:p>' %
                      (nsdecls('a'), self.test_text))
        self.p = parse_xml_bytes(self.p_xml)
        self.paragraph = _Paragraph(self.p)

    def test_runs_size(self):
        """_Paragraph.runs is expected size"""
        # setup ------------------------
        actual_lengths = []
        for p in self.pList:
            paragraph = _Paragraph(p)
            # exercise ----------------
            actual_lengths.append(len(paragraph.runs))
        # verify ------------------
        expected = [0, 0, 2, 1, 1, 1]
        actual = actual_lengths
        msg = "expected run count %s, got %s" % (expected, actual)
        self.assertEqual(expected, actual, msg)

    def test_add_run_increments_run_count(self):
        """_Paragraph.add_run() increments run count"""
        # setup ------------------------
        p_elm = self.pList[0]
        paragraph = _Paragraph(p_elm)
        # exercise ---------------------
        paragraph.add_run()
        # verify -----------------------
        expected = 1
        actual = len(paragraph.runs)
        msg = "expected run count %s, got %s" % (expected, actual)
        self.assertEqual(expected, actual, msg)

    @patch('pptx.text.ParagraphAlignment')
    def test_alignment_value(self, ParagraphAlignment):
        """_Paragraph.alignment value is calculated correctly"""
        # setup ------------------------
        paragraph = test_text_objects.paragraph
        paragraph._Paragraph__p = __p = MagicMock(name='__p')
        __p.get_algn = get_algn = Mock(name='get_algn')
        get_algn.return_value = algn_val = Mock(name='algn_val')
        alignment = Mock(name='alignment')
        from_text_align_type = ParagraphAlignment.from_text_align_type
        from_text_align_type.return_value = alignment
        # exercise ---------------------
        retval = paragraph.alignment
        # verify -----------------------
        get_algn.assert_called_once_with()
        from_text_align_type.assert_called_once_with(algn_val)
        assert_that(retval, is_(same_instance(alignment)))

    @patch('pptx.text.ParagraphAlignment')
    def test_alignment_assignment(self, ParagraphAlignment):
        """Assignment to _Paragraph.alignment assigns value"""
        # setup ------------------------
        paragraph = test_text_objects.paragraph
        paragraph._Paragraph__p = __p = MagicMock(name='__p')
        __p.set_algn = set_algn = Mock(name='set_algn')
        algn_val = Mock(name='algn_val')
        to_text_align_type = ParagraphAlignment.to_text_align_type
        to_text_align_type.return_value = algn_val
        alignment = PP.ALIGN_CENTER
        # exercise ---------------------
        paragraph.alignment = alignment
        # verify -----------------------
        to_text_align_type.assert_called_once_with(alignment)
        set_algn.assert_called_once_with(algn_val)

    def test_alignment_integrates_with_CT_TextParagraph(self):
        """_Paragraph.alignment integrates with CT_TextParagraph"""
        # setup ------------------------
        paragraph = test_text_objects.paragraph
        expected_xml = test_text_xml.centered_paragraph
        # exercise ---------------------
        paragraph.alignment = PP.ALIGN_CENTER
        # verify -----------------------
        self.assertEqualLineByLine(expected_xml, paragraph._Paragraph__p)

    def test_clear_removes_all_runs(self):
        """_Paragraph.clear() removes all runs from paragraph"""
        # setup ------------------------
        p = self.pList[2]
        SubElement(p, 'a:pPr')
        paragraph = _Paragraph(p)
        assert_that(len(paragraph.runs), is_(equal_to(2)))
        # exercise ---------------------
        paragraph.clear()
        # verify -----------------------
        assert_that(len(paragraph.runs), is_(equal_to(0)))

    def test_clear_preserves_paragraph_properties(self):
        """_Paragraph.clear() preserves paragraph properties"""
        # setup ------------------------
        p_xml = ('<a:p %s><a:pPr lvl="1"/><a:r><a:t>%s</a:t></a:r></a:p>' %
                 (nsdecls('a'), self.test_text))
        p_elm = parse_xml_bytes(p_xml)
        paragraph = _Paragraph(p_elm)
        expected_p_xml = '<a:p %s><a:pPr lvl="1"/></a:p>' % nsdecls('a')
        # exercise ---------------------
        paragraph.clear()
        # verify -----------------------
        p_xml = serialize_xml(paragraph._Paragraph__p)
        assert_that(p_xml, is_(equal_to(expected_p_xml)))

    def test_level_setter_generates_correct_xml(self):
        """_Paragraph.level setter generates correct XML"""
        # setup ------------------------
        expected_xml = (
            '<a:p %s>\n  <a:pPr lvl="2"/>\n  <a:r>\n    <a:t>test text</a:t>'
            '\n  </a:r>\n</a:p>\n' % nsdecls('a')
        )
        # exercise ---------------------
        self.paragraph.level = 2
        # verify -----------------------
        self.assertEqualLineByLine(expected_xml, self.paragraph._Paragraph__p)

    def test_level_default_is_zero(self):
        """_Paragraph.level defaults to zero on no lvl attribute"""
        # verify -----------------------
        assert_that(self.paragraph.level, is_(equal_to(0)))

    def test_level_roundtrips_intact(self):
        """_Paragraph.level property round-trips intact"""
        # exercise ---------------------
        self.paragraph.level = 5
        # verify -----------------------
        assert_that(self.paragraph.level, is_(equal_to(5)))

    def test_level_raises_on_bad_value(self):
        """_Paragraph.level raises on attempt to assign invalid value"""
        test_cases = ('0', -1, 9)
        for value in test_cases:
            with self.assertRaises(ValueError):
                self.paragraph.level = value

    def test_set_font_size(self):
        """Assignment to _Paragraph.font.size changes font size"""
        # setup ------------------------
        newfontsize = Pt(54.3)
        expected_xml = (
            '<a:p %s>\n  <a:pPr>\n    <a:defRPr sz="5430"/>\n  </a:pPr>\n  <a'
            ':r>\n    <a:t>test text</a:t>\n  </a:r>\n</a:p>\n' % nsdecls('a')
        )
        # exercise ---------------------
        self.paragraph.font.size = newfontsize
        # verify -----------------------
        self.assertEqualLineByLine(expected_xml, self.paragraph._Paragraph__p)

    def test_text_setter_sets_single_run_text(self):
        """assignment to _Paragraph.text creates single run containing value"""
        # setup ------------------------
        test_text = 'python-pptx was here!!'
        p_elm = self.pList[2]
        paragraph = _Paragraph(p_elm)
        # exercise ---------------------
        paragraph.text = test_text
        # verify -----------------------
        assert_that(len(paragraph.runs), is_(equal_to(1)))
        assert_that(paragraph.runs[0].text, is_(equal_to(test_text)))

    def test_text_accepts_non_ascii_strings(self):
        """assignment of non-ASCII string to text does not raise"""
        # setup ------------------------
        _7bit_string = 'String containing only 7-bit (ASCII) characters'
        _8bit_string = '8-bit string: Hér er texti með íslenskum stöfum.'
        _utf8_literal = u'unicode literal: Hér er texti með íslenskum stöfum.'
        _utf8_from_8bit = unicode('utf-8 unicode: Hér er texti', 'utf-8')
        # verify -----------------------
        try:
            text = _7bit_string
            self.paragraph.text = text
            text = _8bit_string
            self.paragraph.text = text
            text = _utf8_literal
            self.paragraph.text = text
            text = _utf8_from_8bit
            self.paragraph.text = text
        except ValueError:
            msg = "_Paragraph.text rejects valid text string '%s'" % text
            self.fail(msg)


class Describe_Run(object):

    def it_can_get_the_text_of_the_run(self, run, test_text):
        assert run.text == test_text

    def it_can_change_the_text_of_the_run(self, run):
        run.text = 'new text'
        assert run.text == 'new text'

    # fixtures ---------------------------------------------

    @pytest.fixture
    def test_text(self):
        return 'test text'

    @pytest.fixture
    def r_xml(self, test_text):
        return ('<a:r %s><a:t>%s</a:t></a:r>' %
                (nsdecls('a'), test_text))

    @pytest.fixture
    def r(self, r_xml):
        return parse_xml_bytes(r_xml)

    @pytest.fixture
    def run(self, r):
        return _Run(r)


class DescribeTextFrame(TestCase):

    def setUp(self):
        path = absjoin(test_file_dir, 'slide1.xml')
        self.sld = parse_xml_file(path).getroot()
        xpath = './p:cSld/p:spTree/p:sp/p:txBody'
        self.txBodyList = self.sld.xpath(xpath, namespaces=nsmap)

    def test_paragraphs_size(self):
        """TextFrame.paragraphs is expected size"""
        # setup ------------------------
        actual_lengths = []
        for txBody in self.txBodyList:
            textframe = TextFrame(txBody)
            # exercise ----------------
            actual_lengths.append(len(textframe.paragraphs))
        # verify -----------------------
        expected = [1, 1, 2, 1, 1]
        actual = actual_lengths
        msg = "expected paragraph count %s, got %s" % (expected, actual)
        self.assertEqual(expected, actual, msg)

    def test_add_paragraph_xml(self):
        """TextFrame.add_paragraph does what it says"""
        # setup ------------------------
        txBody_xml = (
            '<p:txBody %s><a:bodyPr/><a:p><a:r><a:t>Test text</a:t></a:r></a:'
            'p></p:txBody>' % nsdecls('p', 'a')
        )
        expected_xml = (
            '<p:txBody %s><a:bodyPr/><a:p><a:r><a:t>Test text</a:t></a:r></a:'
            'p><a:p/></p:txBody>' % nsdecls('p', 'a')
        )
        txBody = parse_xml_bytes(txBody_xml)
        textframe = TextFrame(txBody)
        # exercise ---------------------
        textframe.add_paragraph()
        # verify -----------------------
        assert_that(len(textframe.paragraphs), is_(equal_to(2)))
        textframe_xml = serialize_xml(textframe._txBody)
        expected = expected_xml
        actual = textframe_xml
        msg = "\nExpected: '%s'\n\n     Got: '%s'" % (expected, actual)
        if not expected == actual:
            raise AssertionError(msg)

    def test_text_setter_structure_and_value(self):
        """Assignment to TextFrame.text yields single run para set to value"""
        # setup ------------------------
        test_text = 'python-pptx was here!!'
        txBody = self.txBodyList[2]
        textframe = TextFrame(txBody)
        # exercise ---------------------
        textframe.text = test_text
        # verify paragraph count -------
        expected = 1
        actual = len(textframe.paragraphs)
        msg = "expected paragraph count %s, got %s" % (expected, actual)
        self.assertEqual(expected, actual, msg)
        # verify value -----------------
        expected = test_text
        actual = textframe.paragraphs[0].runs[0].text
        msg = "expected text '%s', got '%s'" % (expected, actual)
        self.assertEqual(expected, actual, msg)

    def test_vertical_anchor_works(self):
        """Assignment to TextFrame.vertical_anchor sets vert anchor"""
        # setup ------------------------
        txBody_xml = (
            '<p:txBody %s><a:bodyPr/><a:p><a:r><a:t>Test text</a:t></a:r></a:'
            'p></p:txBody>' % nsdecls('p', 'a')
        )
        expected_xml = (
            '<p:txBody %s>\n  <a:bodyPr anchor="ctr"/>\n  <a:p>\n    <a:r>\n '
            '     <a:t>Test text</a:t>\n    </a:r>\n  </a:p>\n</p:txBody>\n' %
            nsdecls('p', 'a')
        )
        txBody = parse_xml_bytes(txBody_xml)
        textframe = TextFrame(txBody)
        # exercise ---------------------
        textframe.vertical_anchor = MSO.ANCHOR_MIDDLE
        # verify -----------------------
        self.assertEqualLineByLine(expected_xml, textframe._txBody)

    def test_word_wrap_works(self):
        """Assignment to TextFrame.word_wrap sets word wrap value"""
        # setup ------------------------
        txBody_xml = (
            '<p:txBody %s><a:bodyPr/><a:p><a:r><a:t>Test text</a:t></a:r></a:'
            'p></p:txBody>' % nsdecls('p', 'a')
        )
        true_expected_xml = (
            '<p:txBody %s>\n  <a:bodyPr wrap="square"/>\n  <a:p>\n    <a:r>\n '
            '     <a:t>Test text</a:t>\n    </a:r>\n  </a:p>\n</p:txBody>\n' %
            nsdecls('p', 'a')
        )
        false_expected_xml = (
            '<p:txBody %s>\n  <a:bodyPr wrap="none"/>\n  <a:p>\n    <a:r>\n '
            '     <a:t>Test text</a:t>\n    </a:r>\n  </a:p>\n</p:txBody>\n' %
            nsdecls('p', 'a')
        )
        none_expected_xml = (
            '<p:txBody %s>\n  <a:bodyPr/>\n  <a:p>\n    <a:r>\n '
            '     <a:t>Test text</a:t>\n    </a:r>\n  </a:p>\n</p:txBody>\n' %
            nsdecls('p', 'a')
        )

        txBody = parse_xml_bytes(txBody_xml)
        textframe = TextFrame(txBody)

        self.assertEqual(textframe.word_wrap, None)

        # exercise ---------------------
        textframe.word_wrap = True
        # verify -----------------------
        self.assertEqualLineByLine(
            true_expected_xml, textframe._txBody)
        self.assertEqual(textframe.word_wrap, True)

        # exercise ---------------------
        textframe.word_wrap = False
        # verify -----------------------
        self.assertEqualLineByLine(
            false_expected_xml, textframe._txBody)
        self.assertEqual(textframe.word_wrap, False)

        # exercise ---------------------
        textframe.word_wrap = None
        # verify -----------------------
        self.assertEqualLineByLine(
            none_expected_xml, textframe._txBody)
        self.assertEqual(textframe.word_wrap, None)
